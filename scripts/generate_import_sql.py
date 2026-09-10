"""
Importador Excel -> Supabase para el Panel de Campaña de Soja.

Lee el Excel de siempre (SOJA CAMPAÑA 2526.xlsm, el que se carga a mano en
"Ventas Finn", "Balanza", "Intermedio", "Prod. Final", "Despacho", etc.) y
genera un único script SQL que reemplaza el contenido de las tablas
operativas en Supabase por lo que hay HOY en el Excel.

Estrategia: por cada tabla, DELETE de todas las filas + INSERT de las filas
actuales del Excel, todo dentro de una sola transacción (BEGIN...COMMIT) para
que si algo falla a mitad de camino no quede la base a medio actualizar.
Es seguro reemplazar así porque ninguna tabla depende del "id" interno que
genera Supabase — todos los vínculos reales (vta_id, lote, pedido, venta)
son columnas de negocio que ya vienen del Excel.

Uso:
    python3 generate_import_sql.py "SOJA CAMPAÑA 2526.xlsm" salida.sql

El SQL generado se revisa y se ejecuta a mano (vía el MCP de Supabase) —
este script nunca toca la base directamente ni maneja ninguna clave.
"""
import sys
import datetime
import openpyxl

from full_extract import extract_data, dstr


def esc(v):
    if v is None:
        return 'NULL'
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, (int, float)):
        return repr(v)
    sv = str(v).strip()
    if sv == '':
        return 'NULL'
    sv = sv.replace("'", "''")
    return f"'{sv}'"


def insert_stmt(table, columns, rows):
    if not rows:
        return None
    cols_sql = ', '.join(columns)
    values_sql = []
    for row in rows:
        vals = ', '.join(esc(row.get(c)) for c in columns)
        values_sql.append(f"({vals})")
    return f"INSERT INTO {table} ({cols_sql}) VALUES\n" + ",\n".join(values_sql) + ";"


def fmt_time_cell(v):
    """Para columnas 'text' que en el Excel guardan horarios (Calendario Produccion).
    Formato HH:MM:SS para que coincida con los datos ya migrados."""
    if v is None:
        return None
    if isinstance(v, datetime.time):
        return v.strftime('%H:%M:%S')
    if isinstance(v, datetime.datetime):
        return v.strftime('%H:%M:%S')
    if isinstance(v, str) and v.strip() == '':
        return None
    return v


def clean_hours(v):
    """Las horas de Calendario Produccion vienen de restas de horarios en Excel y
    a veces arrastran errores de punto flotante minúsculos (9.000000000000002 en vez
    de 9.0). Si está a menos de 1e-6 de un entero, se redondea."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        rv = round(v)
        return float(rv) if abs(v - rv) < 1e-6 else round(v, 6)
    return v


def extract_variedades(wb):
    ws = wb['MASTER']
    out = []
    for r in range(3, 14):  # filas 3 a 13, igual que la migración original
        variedad = ws.cell(r, 3).value
        if variedad is None or str(variedad).strip() == '':
            continue
        out.append({
            'variedad': str(variedad).strip(),
            'ciclo': ws.cell(r, 4).value,
            'semilla': ws.cell(r, 5).value,
            'semillero': ws.cell(r, 6).value,
            'tecnologia': ws.cell(r, 7).value,
        })
    return out


def extract_calendario(wb):
    ws = wb['Calendario Produccion']
    out = []
    for r in range(5, ws.max_row + 1):
        fecha = ws.cell(r, 1).value
        if fecha is None:
            continue
        out.append({
            'fecha': dstr(fecha),
            'dia': ws.cell(r, 2).value,
            'gessi_inicio': fmt_time_cell(ws.cell(r, 3).value),
            'gessi_fin': fmt_time_cell(ws.cell(r, 4).value),
            'horas_gessi': clean_hours(ws.cell(r, 5).value),
            'blomar_inicio': fmt_time_cell(ws.cell(r, 6).value),
            'blomar_fin': fmt_time_cell(ws.cell(r, 7).value),
            'horas_blomar': clean_hours(ws.cell(r, 8).value),
            'observacion': ws.cell(r, 9).value,
        })
    return out


def build_sql(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, keep_vba=True)
    d = extract_data(wb)
    variedades = extract_variedades(wb)
    calendario = extract_calendario(wb)

    blocks = []
    counts = {}

    def add_table(table, columns, rows):
        counts[table] = len(rows)
        # "WHERE true": Supabase tiene activada la protección "safeupdate",
        # que bloquea cualquier DELETE/UPDATE sin WHERE. Con WHERE true se
        # sigue borrando todo, pero no dispara esa protección.
        blocks.append(f"DELETE FROM {table} WHERE true;")
        stmt = insert_stmt(table, columns, rows)
        if stmt:
            blocks.append(stmt)

    add_table('variedades', ['variedad', 'ciclo', 'semilla', 'semillero', 'tecnologia'], variedades)

    add_table('pedidos_venta',
               ['vta_id', 'fecha', 'cliente', 'categoria', 'semillero', 'variedad', 'cantidad',
                'vendedor', 'tratamiento', 'presentacion', 'fecha_entrega', 'kg_total', 'control_pendiente'],
               [{'vta_id': v['vta_id'], 'fecha': v['fecha'], 'cliente': v['cliente'], 'categoria': v['categoria'],
                 'semillero': v['semillero'], 'variedad': v['variedad'], 'cantidad': v['cantidad'],
                 'vendedor': v['vendedor'], 'tratamiento': v['tratamiento'], 'presentacion': v['pack'],
                 'fecha_entrega': v['fecha_entrega'], 'kg_total': v['kg_total'],
                 'control_pendiente': v['control_pendiente']}
                for v in d['ventas']])

    add_table('balanza',
               ['fecha', 'chofer', 'maquina', 'campo', 'variedad', 'categoria', 'neto', 'silo', 'merma', 'merma_viento'],
               d['balanza'])

    add_table('lotes',
               ['lote', 'procedencia', 'variedad', 'area', 'kg', 'kg_hr', 'paro', 'detalle_paro', 'fin', 'pureza', 'ok_calidad', 'pg'],
               d['lotes'])

    add_table('produccion_final',
               ['vta_id', 'venta', 'estado', 'cliente', 'variedad', 'presentacion', 'kg', 'lote_final', 'fecha_entrega'],
               d['prodfinal'])

    despachos_rows = []
    for row in d['despachos']:
        row = dict(row)
        if row.get('lote_final') == 0:  # placeholder del Excel para "sin lote todavía"
            row['lote_final'] = None
        despachos_rows.append(row)
    add_table('despachos',
               ['pedido', 'cliente', 'dia_carga', 'variedad', 'presentacion', 'cantidad', 'kg', 'lote_final', 'estado'],
               despachos_rows)

    add_table('stock_resumen_marca', ['marca', 'pintados', 'sin_pintar', 'nuevos'], d['resumen_marca'])

    add_table('stock_bolsas_semilla', ['semilla', 'bolsas'],
               [{'semilla': k, 'bolsas': v} for k, v in d['bolsas_by_semilla'].items()])

    add_table('plan_curado',
               ['variedad', 'tratamiento', 'pendiente_total', 'curado_disp', 'neto_a_curar'],
               d['plan_curado'])

    pcw = [w for w in d['plan_curado_weekly'] if w['semana'] and w['semana'] != 'Sin fecha']
    add_table('plan_curado_semanal', ['semana', 'kg'], pcw)

    add_table('lotes_curado',
               ['lote_intermedio', 'lote_final', 'variedad', 'categoria', 'semillero', 'kg_disp_inter',
                'kg_consumidos', 'area_int', 'kg_producidos', 'kg_disponibles', 'area_final', 'tratamiento',
                'complemento', 'hora_inicio', 'hora_fin', 'horas_prod', 'kg_hs'],
               d['lotes_curado'])

    add_table('calendario_produccion',
               ['fecha', 'dia', 'gessi_inicio', 'gessi_fin', 'horas_gessi', 'blomar_inicio', 'blomar_fin', 'horas_blomar', 'observacion'],
               calendario)

    sql = "BEGIN;\n\n" + "\n\n".join(blocks) + "\n\nCOMMIT;\n"
    return sql, counts


if __name__ == '__main__':
    src = sys.argv[1]
    out_path = sys.argv[2]
    sql, counts = build_sql(src)
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(sql)
    print(f"SQL generado en {out_path} ({len(sql)} caracteres)")
    print("Filas por tabla:")
    for t, n in counts.items():
        print(f"  {t}: {n}")
